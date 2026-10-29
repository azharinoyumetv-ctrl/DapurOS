from permissions import require_permission
from paging import encode_cursor, decode_cursor
import re
"""Product CRUD + bulk Excel/CSV import."""
import io
import uuid
from typing import List, Optional
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Query, Header

from database import get_db, utcnow
from models import ProductCreate, ProductUpdate, BulkImportResult
from auth import get_current_user

router = APIRouter(prefix="/api/products", tags=["products"])


def get_module(x_dagangos_module: str = Header(default="dapuros", alias="X-DagangOS-Module")):
    """Modul aplikasi yang meminta (dapuros / geraina). Untuk memisahkan katalog per aplikasi."""
    return (x_dagangos_module or "dapuros").lower()


@router.get('/catalogue', dependencies=[Depends(require_permission('products.read'))])
async def catalogue_page(user: dict = Depends(get_current_user), cursor: Optional[str] = None,
                         limit: int = Query(500, ge=1, le=500), active_only: bool = False):
    db = get_db()
    scope = {'store_id': user['store_id'], 'module': 'dapuros'}
    if active_only:
        scope['active'] = {'$ne': False}
    query = dict(scope)
    if cursor:
        query['id'] = {'$gt': decode_cursor(cursor, ['id'])['id']}
    rows = await db.products.find(query, {'_id': 0}).sort('id', 1).limit(limit + 1).to_list(length=limit + 1)
    more = len(rows) > limit
    items = rows[:limit]
    return {'items': items, 'total': await db.products.count_documents(scope),
            'next_cursor': encode_cursor({'id': items[-1]['id']}) if more else None}


@router.get('/lookup', dependencies=[Depends(require_permission('products.read'))])
async def lookup_product(code: str = Query(min_length=1, max_length=200), user: dict = Depends(get_current_user)):
    query = {'store_id': user['store_id'], 'module': 'dapuros', 'active': {'$ne': False},
             '$or': [{'id': code}, {'sku': {'$regex': '^' + re.escape(code) + '$', '$options': 'i'}}]}
    rows = await get_db().products.find(query, {'_id': 0}).limit(2).to_list(length=2)
    if not rows:
        raise HTTPException(404, 'Barcode tidak ditemukan dalam katalog toko.')
    if len(rows) > 1:
        raise HTTPException(409, 'Barcode dipakai beberapa produk. Perbaiki SKU produk sebelum memindai.')
    return rows[0]


def _strip_id(p: dict) -> dict:
    p.pop("_id", None)
    return p


@router.get("", dependencies=[Depends(require_permission("products.read"))])
async def list_products(
    user: dict = Depends(get_current_user),
    q: Optional[str] = Query(None),
    category: Optional[str] = Query(None),
    limit: int = 500,
):
    db = get_db()
    flt = {"store_id": user["store_id"]}
    if q:
        flt["$or"] = [
            {"name": {"$regex": q, "$options": "i"}},
            {"sku": {"$regex": q, "$options": "i"}},
            {"category": {"$regex": q, "$options": "i"}},
        ]
    if category and category != "all":
        flt["category"] = category
    cursor = db.products.find(flt, {"_id": 0}).sort("name", 1).limit(limit)
    return await cursor.to_list(length=limit)


@router.get("/category-names", dependencies=[Depends(require_permission("products.read"))])
async def list_category_names(user: dict = Depends(get_current_user)):
    """Daftar nama kategori unik dari produk (untuk filter chip Kasir)."""
    db = get_db()
    cats = await db.products.distinct("category", {"store_id": user["store_id"]})
    return [c for c in cats if c]


@router.post("", dependencies=[Depends(require_permission("products.write"))])
async def create_product(payload: ProductCreate, user: dict = Depends(get_current_user)):
    db=get_db()
    return await transactions.run_transaction(db,lambda session:save_product(db,user,payload.model_dump(),session))


@router.put("/{product_id}", dependencies=[Depends(require_permission("products.write"))])
async def update_product(product_id: str, payload: ProductUpdate, user: dict = Depends(get_current_user)):
    db=get_db()
    values={key:value for key,value in payload.model_dump(exclude_unset=True).items() if value is not None and key!='expected_stock'}
    if not values:
        raise HTTPException(400,'Tidak ada perubahan.')
    return await transactions.run_transaction(db,lambda session:save_product(db,user,values,session,product_id,payload.expected_stock))


@router.delete("/{product_id}", dependencies=[Depends(require_permission("products.write"))])
async def delete_product(product_id: str, user: dict = Depends(get_current_user)):
    db = get_db()
    if await db.stock_transfers.find_one({'store_id':user['store_id'],'status':'Shipped','inventory_applied':True,'items.product_id':product_id}):
        raise HTTPException(409,'Produk sedang dalam pengiriman cabang. Terima atau batalkan transfer sebelum menghapus produk.')
    if await db.purchase_orders.find_one({'store_id':user['store_id'],'status':{'$in':['Ordered','Partial']},'items.product_id':product_id}):
        raise HTTPException(409,'Produk masih ada pada PO yang belum diterima penuh.')

    from bson import ObjectId
    id_or_oid = [product_id]
    try:
        id_or_oid.append(ObjectId(product_id))
    except Exception:
        pass

    res = await db.products.delete_one(
        {"store_id": user["store_id"], "$or": [{"id": {"$in": id_or_oid}}, {"_id": {"$in": id_or_oid}}]}
    )
    if res.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Produk tidak ditemukan")
    return {"ok": True}


@router.post("/bulk-import", response_model=BulkImportResult, dependencies=[Depends(require_permission("products.write"))])
async def bulk_import(
    file: UploadFile = File(...),
    user: dict = Depends(get_current_user),
):
    """Import dari Excel (.xlsx) atau CSV.

    Kolom wajib: name, price
    Kolom opsional: sku, cost, stock, category, unit, description, image_url, active
    """
    db = get_db()
    contents = await file.read(10 * 1024 * 1024 + 1)
    if len(contents) > 10 * 1024 * 1024:
        raise HTTPException(413, "Batas berkas impor adalah 10 MB.")
    fname = (file.filename or "").lower()

    rows_data = []
    try:
        if fname.endswith(".csv"):
            import csv
            try:
                text = contents.decode("utf-8-sig")
            except UnicodeDecodeError:
                text = contents.decode("latin1")

            reader = csv.DictReader(io.StringIO(text))
            if reader.fieldnames:
                field_map = {f: str(f).strip().lower() for f in reader.fieldnames if f is not None}
                for row in reader:
                    row_dict = {}
                    for orig_k, val in row.items():
                        if orig_k in field_map:
                            row_dict[field_map[orig_k]] = val
                    rows_data.append(row_dict)
            else:
                raise Exception("File CSV tidak memiliki header")

        elif fname.endswith(".xlsx") or fname.endswith(".xls"):
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(contents), data_only=True)
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
            if not rows or len(rows) < 1:
                raise Exception("File Excel kosong")

            header = []
            for cell in rows[0]:
                if cell is not None:
                    header.append(str(cell).strip().lower())
                else:
                    header.append("")

            header_map = {name: i for i, name in enumerate(header) if name}

            for row in rows[1:]:
                if all(cell is None for cell in row):
                    continue
                row_dict = {}
                for name, idx in header_map.items():
                    if idx < len(row):
                        row_dict[name] = row[idx]
                    else:
                        row_dict[name] = None
                rows_data.append(row_dict)
        else:
            raise HTTPException(status_code=400, detail="Format file harus .csv, .xls, atau .xlsx")
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Gagal membaca file: {e}")

    # Check headers
    headers_parsed = []
    if fname.endswith(".csv") and 'field_map' in locals():
        headers_parsed = list(field_map.values())
    elif (fname.endswith(".xlsx") or fname.endswith(".xls")) and 'header_map' in locals():
        headers_parsed = list(header_map.keys())

    if "name" not in headers_parsed or "price" not in headers_parsed:
        raise HTTPException(status_code=400, detail="Kolom 'name' dan 'price' wajib ada")

    inserted = updated = skipped = 0
    errors: List[str] = []

    def _is_not_empty(val) -> bool:
        if val is None:
            return False
        s = str(val).strip()
        if s == "" or s.lower() == "nan" or s.lower() == "null":
            return False
        return True

    if len(rows_data)>10000:
        raise HTTPException(422,"Maksimal 10.000 baris per impor. Bagi berkas menjadi beberapa bagian.")
    for idx, row in enumerate(rows_data):
        try:
            name = str(row.get("name") or "").strip()
            if not name or name.lower() == "nan":
                skipped += 1
                continue

            price_val = row.get("price")
            if not _is_not_empty(price_val):
                errors.append(f"Baris {idx + 2}: Kolom 'price' wajib diisi")
                continue
            try:
                price = float(price_val)
            except (TypeError, ValueError):
                errors.append(f"Baris {idx + 2}: Nilai 'price' tidak valid ({price_val!r})")
                continue

            sku = str(row.get("sku")).strip() if _is_not_empty(row.get("sku")) else None

            cost_val = row.get("cost")
            cost = float(cost_val) if _is_not_empty(cost_val) else 0.0

            stock_val = row.get("stock")
            stock_number = float(stock_val) if _is_not_empty(stock_val) else 0.0
            if not math.isfinite(stock_number) or not stock_number.is_integer():
                raise ValueError('Stok harus berupa bilangan bulat yang valid.')
            stock = int(stock_number)

            category = str(row.get("category")).strip() if _is_not_empty(row.get("category")) else "Umum"
            unit = str(row.get("unit")).strip() if _is_not_empty(row.get("unit")) else "pcs"
            description = str(row.get("description")).strip() if _is_not_empty(row.get("description")) else None
            image_url = str(row.get("image_url")).strip() if _is_not_empty(row.get("image_url")) else None

            active_val = row.get("active")
            if _is_not_empty(active_val):
                active_str = str(active_val).lower().strip()
                active = active_str not in ("false", "0", "no", "f", "n")
            else:
                active = True

            doc = {
                "name": name,
                "sku": sku,
                "price": price,
                "cost": cost,
                "stock": stock,
                "category": category,
                "unit": unit,
                "description": description,
                "image_url": image_url,
                "active": active,
                "store_id": user["store_id"],
                "updated_at": utcnow().isoformat(),
            }

            validated = ProductCreate.model_validate(doc).model_dump()
            # Upsert by SKU within store, else insert new
            existing = None
            if doc["sku"]:
                existing = await db.products.find_one({"store_id": user["store_id"], "sku": doc["sku"]})

            if existing:
                await transactions.run_transaction(db,lambda session:save_product(db,user,validated,session,existing['id'],existing.get('stock',0)))
                updated += 1
            else:
                await transactions.run_transaction(db,lambda session:save_product(db,user,validated,session))
                inserted += 1
        except Exception as e:
            errors.append(f"Baris {idx + 2}: {e}")

    return BulkImportResult(inserted=inserted, updated=updated, skipped=skipped, errors=errors[:20])


@router.get("/import-template.csv", dependencies=[Depends(require_permission("products.write"))])
async def import_template():
    """Plain CSV template for bulk import."""
    from fastapi.responses import PlainTextResponse
    text = "name,sku,price,cost,stock,category,unit,description,active\n"
    text += "Es Kopi Susu Gula Aren,GR-CFE-001,22000,9000,100,Minuman,cup,Signature kopi susu,true\n"
    text += "Croissant Coklat,GR-BRD-002,20000,8000,30,Roti,pcs,Croissant isi coklat belgia,true\n"
    return PlainTextResponse(text, media_type="text/csv")

def _strip_id(p: dict) -> dict:
    p.pop("_id", None)
    return p


@router.get("", dependencies=[Depends(require_permission('products.read'))])
async def list_products(
    user: dict = Depends(get_current_user),
    module: str = Depends(get_module),
    q: Optional[str] = Query(None),
    category: Optional[str] = Query(None),
    limit: int = 500,
):
    db = get_db()
    flt = {"store_id": user["store_id"], "module": module}
    if q:
        flt["$or"] = [
            {"name": {"$regex": q, "$options": "i"}},
            {"sku": {"$regex": q, "$options": "i"}},
            {"category": {"$regex": q, "$options": "i"}},
        ]
    if category and category != "all":
        flt["category"] = category
    cursor = db.products.find(flt, {"_id": 0}).sort("name", 1).limit(limit)
    return await cursor.to_list(length=limit)


@router.get("/category-names", dependencies=[Depends(require_permission('products.read'))])
async def list_category_names(user: dict = Depends(get_current_user), module: str = Depends(get_module)):
    """Daftar nama kategori unik dari produk (untuk filter chip Kasir)."""
    db = get_db()
    cats = await db.products.distinct("category", {"store_id": user["store_id"], "module": module})
    return [c for c in cats if c]


@router.post("", dependencies=[Depends(require_permission('products.write'))])
async def create_product(payload: ProductCreate, user: dict = Depends(get_current_user), module: str = Depends(get_module)):
    db = get_db()
    from plan_limits import check_capacity
    await check_capacity(db, user["store_id"], user.get("plan"), "products", "max_products")
    doc = {
        "id": str(uuid.uuid4()),
        "store_id": user["store_id"],
        "module": module,
        "created_at": utcnow().isoformat(),
        "updated_at": utcnow().isoformat(),
        **payload.model_dump(),
    }
    await db.products.insert_one(doc)
    return _strip_id(doc)


@router.put("/{product_id}", dependencies=[Depends(require_permission('products.write'))])
async def update_product(product_id: str, payload: ProductUpdate, user: dict = Depends(get_current_user)):
    db = get_db()
    data = payload.model_dump(exclude_unset=True)
    update = {k: v for k, v in data.items() if v is not None}
    # Izinkan menghapus foto produk secara eksplisit (image_url = null)
    if "image_url" in data and data["image_url"] is None:
        update["image_url"] = None
    if not update:
        raise HTTPException(status_code=400, detail="Tidak ada perubahan")
    update["updated_at"] = utcnow().isoformat()

    from bson import ObjectId
    id_or_oid = [product_id]
    try:
        id_or_oid.append(ObjectId(product_id))
    except Exception:
        pass

    res = await db.products.find_one_and_update(
        {"store_id": user["store_id"], "$or": [{"id": {"$in": id_or_oid}}, {"_id": {"$in": id_or_oid}}]},
        {"$set": update},
        return_document=True,
        projection={"_id": 0},
    )
    if not res:
        raise HTTPException(status_code=404, detail="Produk tidak ditemukan")
    return res


@router.delete("/{product_id}", dependencies=[Depends(require_permission('products.write'))])
async def delete_product(product_id: str, user: dict = Depends(get_current_user)):
    db = get_db()

    from bson import ObjectId
    id_or_oid = [product_id]
    try:
        id_or_oid.append(ObjectId(product_id))
    except Exception:
        pass

    res = await db.products.delete_one(
        {"store_id": user["store_id"], "$or": [{"id": {"$in": id_or_oid}}, {"_id": {"$in": id_or_oid}}]}
    )
    if res.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Produk tidak ditemukan")
    return {"ok": True}


@router.post("/bulk-import", response_model=BulkImportResult, dependencies=[Depends(require_permission('products.write'))])
async def bulk_import(
    file: UploadFile = File(...),
    user: dict = Depends(get_current_user),
    module: str = Depends(get_module),
):
    """Import dari Excel (.xlsx) atau CSV.

    Kolom wajib: name, price
    Kolom opsional: sku, cost, stock, category, unit, description, image_url, active
    """
    db = get_db()
    contents = await file.read()
    fname = (file.filename or "").lower()

    rows_data = []
    try:
        if fname.endswith(".csv"):
            import csv
            try:
                text = contents.decode("utf-8-sig")
            except UnicodeDecodeError:
                text = contents.decode("latin1")

            reader = csv.DictReader(io.StringIO(text))
            if reader.fieldnames:
                field_map = {f: str(f).strip().lower() for f in reader.fieldnames if f is not None}
                for row in reader:
                    row_dict = {}
                    for orig_k, val in row.items():
                        if orig_k in field_map:
                            row_dict[field_map[orig_k]] = val
                    rows_data.append(row_dict)
            else:
                raise Exception("File CSV tidak memiliki header")

        elif fname.endswith(".xlsx") or fname.endswith(".xls"):
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(contents), data_only=True)
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
            if not rows or len(rows) < 1:
                raise Exception("File Excel kosong")

            header = []
            for cell in rows[0]:
                if cell is not None:
                    header.append(str(cell).strip().lower())
                else:
                    header.append("")

            header_map = {name: i for i, name in enumerate(header) if name}

            for row in rows[1:]:
                if all(cell is None for cell in row):
                    continue
                row_dict = {}
                for name, idx in header_map.items():
                    if idx < len(row):
                        row_dict[name] = row[idx]
                    else:
                        row_dict[name] = None
                rows_data.append(row_dict)
        else:
            raise HTTPException(status_code=400, detail="Format file harus .csv, .xls, atau .xlsx")
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Gagal membaca file: {e}")

    # Check headers
    headers_parsed = []
    if fname.endswith(".csv") and 'field_map' in locals():
        headers_parsed = list(field_map.values())
    elif (fname.endswith(".xlsx") or fname.endswith(".xls")) and 'header_map' in locals():
        headers_parsed = list(header_map.keys())

    if "name" not in headers_parsed or "price" not in headers_parsed:
        raise HTTPException(status_code=400, detail="Kolom 'name' dan 'price' wajib ada")

    inserted = updated = skipped = 0
    errors: List[str] = []

    def _is_not_empty(val) -> bool:
        if val is None:
            return False
        s = str(val).strip()
        if s == "" or s.lower() == "nan" or s.lower() == "null":
            return False
        return True

    for idx, row in enumerate(rows_data):
        try:
            name = str(row.get("name") or "").strip()
            if not name or name.lower() == "nan":
                skipped += 1
                continue

            price_val = row.get("price")
            if not _is_not_empty(price_val):
                # price is a required column (see docstring) -- silently defaulting a missing
                # price to 0.0 let a row through as a free product with no signal anything was
                # wrong. Same bug class already fixed in GerainaOS (commit 436fbce): treat a
                # missing price as a row-level error instead of writing a broken product.
                errors.append(f"Baris {idx + 2}: kolom 'price' wajib diisi")
                continue
            price = float(price_val)

            sku = str(row.get("sku")).strip() if _is_not_empty(row.get("sku")) else None

            cost_val = row.get("cost")
            cost = float(cost_val) if _is_not_empty(cost_val) else 0.0

            stock_val = row.get("stock")
            stock = int(float(stock_val)) if _is_not_empty(stock_val) else 0

            category = str(row.get("category")).strip() if _is_not_empty(row.get("category")) else "Umum"
            unit = str(row.get("unit")).strip() if _is_not_empty(row.get("unit")) else "pcs"
            description = str(row.get("description")).strip() if _is_not_empty(row.get("description")) else None
            image_url = str(row.get("image_url")).strip() if _is_not_empty(row.get("image_url")) else None

            active_val = row.get("active")
            if _is_not_empty(active_val):
                active_str = str(active_val).lower().strip()
                active = active_str not in ("false", "0", "no", "f", "n")
            else:
                active = True

            doc = {
                "name": name,
                "sku": sku,
                "price": price,
                "cost": cost,
                "stock": stock,
                "category": category,
                "unit": unit,
                "description": description,
                "image_url": image_url,
                "active": active,
                "store_id": user["store_id"],
                "module": module,
                "updated_at": utcnow().isoformat(),
            }

            # Upsert by SKU within store + module, else insert new
            existing = None
            if doc["sku"]:
                existing = await db.products.find_one({"store_id": user["store_id"], "module": module, "sku": doc["sku"]})

            if existing:
                await db.products.update_one({"id": existing["id"]}, {"$set": doc})
                updated += 1
            else:
                doc["id"] = str(uuid.uuid4())
                doc["created_at"] = utcnow().isoformat()
                await db.products.insert_one(doc)
                inserted += 1
        except Exception as e:
            errors.append(f"Baris {idx + 2}: {e}")

    return BulkImportResult(inserted=inserted, updated=updated, skipped=skipped, errors=errors[:20])


@router.get("/import-template.csv", dependencies=[Depends(require_permission('products.read'))])
async def import_template():
    """Plain CSV template for bulk import."""
    from fastapi.responses import PlainTextResponse
    text = "name,sku,price,cost,stock,category,unit,description,active\n"
    text += "Es Kopi Susu Gula Aren,GR-CFE-001,22000,9000,100,Minuman,cup,Signature kopi susu,true\n"
    text += "Croissant Coklat,GR-BRD-002,20000,8000,30,Roti,pcs,Croissant isi coklat belgia,true\n"
    return PlainTextResponse(text, media_type="text/csv")
